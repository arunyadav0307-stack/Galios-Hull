#!/usr/bin/env python3
"""
Generate all 4 original publication-grade black-and-white figures for the Nova Science Publishers
book chapter in full compliance with Nova Chapter Guidelines:
- Portrait width: 4.40 inches (strictly <= 4.5 inches maximum)
- Resolution: 600 DPI exact (within Nova's 350-750 DPI requirement)
- Color mode: Pure black, white, grayscale fills, and distinct line/marker styles ($0 color fee)
- Separate source formats generated per figure:
  1. Editable native PowerPoint vector file (.pptx) using native shapes/connectors/polylines/textboxes
  2. Standalone editable Word figure document (.docx)
  3. Vector Encapsulated PostScript (.eps)
  4. High-resolution 600 DPI TIFF (.tiff) with embedded 600 DPI header metadata
  5. High-resolution 600 DPI PNG (.png) with embedded 600 DPI header metadata
"""

import os
import math
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle, Polygon, FancyArrowPatch, Arc
from PIL import Image
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor as PptRGB

plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["DejaVu Serif", "Times New Roman", "Times"],
    "mathtext.fontset": "stix",
    "font.size": 7.0,
    "axes.edgecolor": "black",
    "axes.linewidth": 0.9,
    "xtick.color": "black",
    "ytick.color": "black",
    "text.color": "black",
})

OUT_DIRS = [
    "/home/user/Galios-Hull/Figures",
    "/home/user/Galios-Hull/Nova_Chapter_Submission_Package/Figures",
]
for d in OUT_DIRS:
    os.makedirs(d, exist_ok=True)

CAPTIONS = {
    1: (
        "Figure 1. Architectural overview of gate-based quantum computation: "
        "(a) geometric Bloch sphere representation of a pure single-qubit state "
        "|ψ(θ, φ)⟩ = cos(θ/2)|0⟩ + exp(iφ)sin(θ/2)|1⟩; "
        "(b) canonical three-qubit quantum circuit illustrating computational-basis initialization, "
        "Hadamard superposition (H), two-qubit CNOT entanglement, parameterized single-qubit rotations "
        "Ry(θ1) and Rz(θ2), and projective measurement onto a classical bit register c3; and "
        "(c) four-tier hardware-to-algorithm execution stack linking physical qubit control pulses to "
        "intelligent computing applications."
    ),
    2: (
        "Figure 2. End-to-end quantum algorithm design, compilation, error suppression, and execution workflow "
        "across noisy intermediate-scale quantum (NISQ) and fault-tolerant computing regimes. Solid arrows "
        "indicate forward synthesis and execution stages; dashed arrows indicate iterative variational parameter "
        "updates and hardware-aware compilation feedback."
    ),
    3: (
        "Figure 3. Closed-loop hybrid quantum–classical intelligent computing architecture. Classical input "
        "features x ∈ R^d undergo preprocessing and quantum state encoding U_Φ(x), followed by trainable "
        "parameterized layers U_var(θ) and shot-based observable estimation on the quantum processing unit (QPU), "
        "while a classical co-processor evaluates empirical loss L(θ) and updates parameters θ^(t+1) via "
        "parameter-shift gradients or derivative-free optimizers."
    ),
    4: (
        "Figure 4. Comparative scaling and regime-selection framework for quantum and classical intelligent "
        "computing: (a) semi-logarithmic comparison of query/gate complexity log10(ops) versus register width "
        "n = log2(N) across unstructured classical search O(N), Grover search O(√N), sub-exponential number field "
        "sieve scaling, and Shor polynomial-time factoring O(n^3); and (b) structured decision matrix classifying "
        "problem regimes across classical baselines, NISQ hybrid heuristics, and fault-tolerant quantum algorithms."
    ),
}


def save_all_raster_and_vector(fig, fig_num):
    """Save figure to 600 DPI PNG, 600 DPI TIFF, vector EPS, and standalone DOCX in all output dirs."""
    tmp_png = f"/tmp/fig_{fig_num}_raw.png"
    fig.savefig(tmp_png, dpi=600, facecolor="white", edgecolor="none")
    im = Image.open(tmp_png).convert("L")

    for out_d in OUT_DIRS:
        png_path = os.path.join(out_d, f"Figure_{fig_num}.png")
        tiff_path = os.path.join(out_d, f"Figure_{fig_num}.tiff")
        eps_path = os.path.join(out_d, f"Figure_{fig_num}.eps")
        docx_path = os.path.join(out_d, f"Figure_{fig_num}.docx")

        im.save(png_path, format="PNG", dpi=(600, 600), optimize=True)
        im.save(tiff_path, format="TIFF", dpi=(600, 600), compression="tiff_lzw")
        fig.savefig(eps_path, format="eps", dpi=600, facecolor="white", edgecolor="none")

        fdoc = docx.Document()
        sec = fdoc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)

        p_img = fdoc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(png_path, width=Inches(4.40))

        p_cap = fdoc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.line_spacing = 1.1
        p_cap.paragraph_format.space_after = Pt(6)
        cap_text = CAPTIONS[fig_num]
        prefix = f"Figure {fig_num}."
        r_bold = p_cap.add_run(prefix)
        r_bold.bold = True
        r_bold.font.name = "Times New Roman"
        r_bold.font.size = Pt(10.0)
        r_rest = p_cap.add_run(cap_text[len(prefix):])
        r_rest.font.name = "Times New Roman"
        r_rest.font.size = Pt(10.0)
        fdoc.save(docx_path)

    plt.close(fig)


def draw_box(ax, x, y, w, h, title, title_size=6.8, fill="#FFFFFF", lw=1.0,
             linestyle="-", rounding=0.04, subtext=None, subsize=5.9, title_y_frac=0.68, sub_y_frac=0.31):
    patch = FancyBboxPatch(
        (x, y), w, h,
        boxstyle=f"round,pad=0.01,rounding_size={rounding}",
        facecolor=fill, edgecolor="black", linewidth=lw, linestyle=linestyle, zorder=3
    )
    ax.add_patch(patch)
    if subtext:
        ax.text(
            x + w / 2, y + h * title_y_frac, title,
            ha="center", va="center", fontsize=title_size, fontweight="bold",
            linespacing=1.12, zorder=4
        )
        ax.text(
            x + w / 2, y + h * sub_y_frac, subtext,
            ha="center", va="center", fontsize=subsize,
            linespacing=1.12, zorder=4
        )
    else:
        ax.text(
            x + w / 2, y + h / 2, title,
            ha="center", va="center", fontsize=title_size, fontweight="bold",
            linespacing=1.15, zorder=4
        )


def arrow(ax, x1, y1, x2, y2, lw=0.95, ls="-"):
    arr = FancyArrowPatch(
        (x1, y1), (x2, y2),
        arrowstyle="-|>,head_length=3.8,head_width=2.3",
        linewidth=lw, linestyle=ls, color="black", zorder=5
    )
    ax.add_patch(arr)


# ==============================================================================
# FIGURE 1: BLOCH SPHERE + 3-QUBIT CIRCUIT + 4-TIER QUANTUM STACK
# ==============================================================================
def build_figure_1():
    fig = plt.figure(figsize=(4.40, 5.25), dpi=600)
    gs = fig.add_gridspec(2, 2, height_ratios=[1.02, 1.15], width_ratios=[0.94, 1.26],
                          hspace=0.18, wspace=0.10, left=0.02, right=0.98, top=0.96, bottom=0.02)

    # Panel (a): Bloch Sphere
    ax_a = fig.add_subplot(gs[0, 0])
    ax_a.set_xlim(-1.52, 1.48)
    ax_a.set_ylim(-1.48, 1.58)
    ax_a.set_aspect("equal")
    ax_a.axis("off")
    ax_a.text(0.0, 1.48, "(a) Qubit Bloch Sphere", fontsize=7.2, fontweight="bold", ha="center")

    R = 0.90
    ax_a.add_patch(Circle((0, 0), R, fill=False, edgecolor="black", lw=1.05))
    th_back = np.linspace(0, np.pi, 100)
    ax_a.plot(R * np.cos(th_back), 0.30 * R * np.sin(th_back), ls="--", lw=0.70, color="#444444")
    th_front = np.linspace(np.pi, 2 * np.pi, 100)
    ax_a.plot(R * np.cos(th_front), 0.30 * R * np.sin(th_front), ls="-", lw=0.85, color="black")

    arrow(ax_a, 0, -1.12, 0, 1.18, lw=0.85)
    arrow(ax_a, 0, 0, 1.16, 0, lw=0.85)
    arrow(ax_a, 0, 0, -0.84, -0.62, lw=0.85)
    ax_a.text(0.07, 1.14, r"$z\;(|0\rangle)$", fontsize=6.9, ha="left", va="bottom")
    ax_a.text(0.07, -1.16, r"$|1\rangle$", fontsize=6.9, ha="left", va="top")
    ax_a.text(1.19, 0.0, r"$y$", fontsize=6.9, ha="left", va="center")
    ax_a.text(-0.90, -0.66, r"$x$", fontsize=6.9, ha="right", va="top")

    vx, vy = 0.52, 0.62
    px, py = 0.52, -0.18
    ax_a.plot([0, px], [0, py], ls=":", lw=0.75, color="black")
    ax_a.plot([px, vx], [py, vy], ls=":", lw=0.75, color="black")
    arrow(ax_a, 0, 0, vx, vy, lw=1.25)
    ax_a.plot(vx, vy, marker="o", ms=3.2, color="black")
    ax_a.text(vx + 0.08, vy + 0.05, r"$|\psi\rangle$", fontsize=7.5, fontweight="bold")
    arc_t = np.linspace(np.pi / 2, np.arctan2(vy, vx), 40)
    ax_a.plot(0.31 * np.cos(arc_t), 0.31 * np.sin(arc_t), lw=0.75, color="black")
    ax_a.text(0.13, 0.36, r"$\theta$", fontsize=6.8)
    arc_p = np.linspace(np.arctan2(-0.62, -0.84), np.arctan2(py, px), 40)
    ax_a.plot(0.28 * np.cos(arc_p), 0.14 * np.sin(arc_p), lw=0.70, color="black")
    ax_a.text(0.10, -0.34, r"$\phi$", fontsize=6.7)

    # Panel (b): 3-Qubit Quantum Circuit
    ax_b = fig.add_subplot(gs[0, 1])
    ax_b.set_xlim(-0.05, 5.25)
    ax_b.set_ylim(0.20, 4.58)
    ax_b.axis("off")
    ax_b.text(2.60, 4.38, "(b) Canonical 3-Qubit Quantum Circuit", fontsize=7.2, fontweight="bold", ha="center")

    ys = [3.62, 2.72, 1.82]
    yc = 0.85
    labels = [r"$|0\rangle_{q_0}$", r"$|0\rangle_{q_1}$", r"$|0\rangle_{q_2}$"]
    for idx, yw in enumerate(ys):
        ax_b.text(0.56, yw, labels[idx], fontsize=7.2, ha="right", va="center")
        ax_b.plot([0.65, 4.30], [yw, yw], lw=0.90, color="black", zorder=1)

    ax_b.text(0.56, yc, r"$c_3$", fontsize=7.2, ha="right", va="center")
    ax_b.plot([0.65, 5.05], [yc + 0.04, yc + 0.04], lw=0.80, color="black")
    ax_b.plot([0.65, 5.05], [yc - 0.04, yc - 0.04], lw=0.80, color="black")
    ax_b.plot([0.92, 1.04], [yc - 0.12, yc + 0.12], lw=0.80, color="black")
    ax_b.text(0.98, yc + 0.21, "3", fontsize=6.2, ha="center")

    draw_box(ax_b, 1.04, ys[0] - 0.26, 0.54, 0.52, r"$H$", title_size=7.4, fill="#FFFFFF", rounding=0.01)
    draw_box(ax_b, 0.92, ys[2] - 0.26, 0.80, 0.52, r"$R_y(\theta_1)$", title_size=6.8, fill="#F2F2F2", rounding=0.01)

    cx1 = 2.10
    ax_b.plot([cx1, cx1], [ys[0], ys[1]], lw=0.95, color="black", zorder=2)
    ax_b.plot(cx1, ys[0], marker="o", ms=4.2, color="black", zorder=4)
    ax_b.add_patch(Circle((cx1, ys[1]), 0.18, facecolor="white", edgecolor="black", lw=0.95, zorder=3))
    ax_b.plot([cx1 - 0.18, cx1 + 0.18], [ys[1], ys[1]], lw=0.90, color="black", zorder=4)
    ax_b.plot([cx1, cx1], [ys[1] - 0.18, ys[1] + 0.18], lw=0.90, color="black", zorder=4)

    cx2 = 2.68
    ax_b.plot([cx2, cx2], [ys[1], ys[2]], lw=0.95, color="black", zorder=2)
    ax_b.plot(cx2, ys[1], marker="o", ms=4.2, color="black", zorder=4)
    ax_b.add_patch(Circle((cx2, ys[2]), 0.18, facecolor="white", edgecolor="black", lw=0.95, zorder=3))
    ax_b.plot([cx2 - 0.18, cx2 + 0.18], [ys[2], ys[2]], lw=0.90, color="black", zorder=4)
    ax_b.plot([cx2, cx2], [ys[2] - 0.18, ys[2] + 0.18], lw=0.90, color="black", zorder=4)

    draw_box(ax_b, 3.06, ys[1] - 0.26, 0.80, 0.52, r"$R_z(\theta_2)$", title_size=6.8, fill="#F2F2F2", rounding=0.01)

    ax_b.plot([4.05, 4.05], [1.48, 3.98], ls="--", lw=0.75, color="#555555")

    mxs = [4.20, 4.39, 4.58]
    for i, yw in enumerate(ys):
        mx = mxs[i]
        ax_b.plot([4.05, mx], [yw, yw], lw=0.90, color="black", zorder=1)
        patch = FancyBboxPatch((mx, yw - 0.24), 0.46, 0.48, boxstyle="square,pad=0.0",
                               facecolor="white", edgecolor="black", lw=0.90, zorder=3)
        ax_b.add_patch(patch)
        arc = Arc((mx + 0.23, yw - 0.09), 0.30, 0.25, angle=0, theta1=0, theta2=180, lw=0.80, zorder=4)
        ax_b.add_patch(arc)
        ax_b.plot([mx + 0.23, mx + 0.36], [yw - 0.09, yw + 0.13], lw=0.85, color="black", zorder=4)
        ax_b.plot([mx + 0.21, mx + 0.21], [yw - 0.24, yc + 0.04], lw=0.60, color="black", zorder=2)
        ax_b.plot([mx + 0.25, mx + 0.25], [yw - 0.24, yc + 0.04], lw=0.60, color="black", zorder=2)
        ax_b.plot(mx + 0.23, yc, marker="v", ms=3.0, color="black", zorder=5)

    # Panel (c): Full-Width 4-Tier Quantum Execution Stack
    ax_c = fig.add_subplot(gs[1, :])
    ax_c.set_xlim(0.0, 10.0)
    ax_c.set_ylim(0.0, 5.55)
    ax_c.axis("off")
    ax_c.text(5.0, 5.28, "(c) Four-Tier Quantum Execution Architecture", fontsize=7.3, fontweight="bold", ha="center")

    tiers = [
        (4.02, "Tier 4: Quantum Algorithms & Intelligent Computing Layer",
         "VQE / QAOA, Quantum Kernels, Grover Search, HHL & QFT Subroutines", "#FFFFFF"),
        (2.74, "Tier 3: Logical Compilation & Hardware Transpilation Layer",
         "Unitary Synthesis, Gate Cancellation, Qubit Routing (SWAP) & Depth Optimization", "#F3F3F3"),
        (1.46, "Tier 2: Error Mitigation & Quantum Error Correction (QEC) Layer",
         "NISQ: ZNE, PEC & Classical Shadows  |  FTQC: Surface Codes & Syndrome Decoding", "#FFFFFF"),
        (0.18, "Tier 1: Physical Quantum Substrate & Pulse Control Layer",
         "Superconducting Transmons, Trapped Ions, Neutral Atoms & Cryogenic RF/Laser Control", "#EAEAEA"),
    ]
    for yt, title, desc, col in tiers:
        draw_box(ax_c, 0.18, yt, 9.64, 0.92, title, title_size=6.7,
                 fill=col, lw=0.95, rounding=0.05, subtext=desc, subsize=5.75)
    for i in range(3):
        y_top = tiers[i][0]
        y_bot = tiers[i + 1][0] + 0.92
        ymid = 0.5 * (y_top + y_bot)
        arrow(ax_c, 3.65, y_top, 3.65, y_bot, lw=0.85)
        arrow(ax_c, 6.35, y_bot, 6.35, y_top, lw=0.85, ls="--")
        if i == 0:
            ax_c.text(3.48, ymid, "Compile", fontsize=5.8, ha="right", va="center")
            ax_c.text(6.52, ymid, "Readout / Diagnostics", fontsize=5.8, ha="left", va="center")

    save_all_raster_and_vector(fig, 1)


# ==============================================================================
# FIGURE 2: END-TO-END QUANTUM ALGORITHM DESIGN & EXECUTION WORKFLOW
# ==============================================================================
def build_figure_2():
    fig, ax = plt.subplots(figsize=(4.40, 5.35), dpi=600)
    fig.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    ax.set_xlim(0.0, 10.0)
    ax.set_ylim(0.0, 12.2)
    ax.axis("off")

    bx, bw = 1.18, 8.62
    cx = bx + bw / 2.0

    draw_box(ax, bx, 10.90, bw, 1.00,
             "1. Computational Problem Specification",
             subtext="Cost Hamiltonian H_C, Oracle f(x), or Dataset {(x_i, y_i)}",
             title_size=6.2, fill="#EDEDED", lw=1.0, rounding=0.08, subsize=5.45)
    arrow(ax, cx, 10.90, cx, 10.35, lw=0.95)

    draw_box(ax, bx, 9.35, bw, 1.00,
             "2. Algorithmic Paradigm & Quantum Encoding Design",
             subtext="Amplitude / Angle / Kernel Map; Variational vs. Oracular Model",
             title_size=6.2, fill="#FFFFFF", lw=1.0, rounding=0.06, subsize=5.45)
    arrow(ax, cx, 9.35, cx, 8.80, lw=0.95)

    draw_box(ax, bx, 7.80, bw, 1.00,
             "3. Hardware-Aware Compilation & Circuit Optimization",
             subtext="Native Gate Synthesis {Rz, √X, CNOT/CZ}, Routing & Depth Reduction",
             title_size=6.2, fill="#F5F5F5", lw=1.0, rounding=0.06, subsize=5.45)
    arrow(ax, cx, 7.80, cx, 7.22, lw=0.95)

    diamond = Polygon([[cx, 7.22], [cx + 2.55, 6.32], [cx, 5.42], [cx - 2.55, 6.32]],
                      closed=True, facecolor="#FFFFFF", edgecolor="black", lw=1.05, zorder=3)
    ax.add_patch(diamond)
    ax.text(cx, 6.48, "Target Hardware", ha="center", va="center", fontsize=6.3, fontweight="bold", zorder=4)
    ax.text(cx, 6.12, "Execution Regime?", ha="center", va="center", fontsize=6.1, zorder=4)

    lx_mid = 3.22
    ax.plot([cx - 2.55, lx_mid], [6.32, 6.32], lw=0.95, color="black")
    arrow(ax, lx_mid, 6.32, lx_mid, 5.12, lw=0.95)
    ax.text(2.75, 6.52, "NISQ", fontsize=6.1, fontweight="bold", ha="right")

    draw_box(ax, 1.18, 3.20, 4.10, 1.90,
             "4a. Pre-Fault-Tolerant\n(NISQ) Execution",
             subtext="Shallow circuits (d << T1, T2);\nZero-Noise Extrapolation (ZNE),\nReadout & Twirling Mitigation",
             title_size=6.1, fill="#FFFFFF", lw=0.95, rounding=0.05, subsize=5.40,
             title_y_frac=0.77, sub_y_frac=0.31)

    rx_mid = 7.75
    ax.plot([cx + 2.55, rx_mid], [6.32, 6.32], lw=0.95, color="black")
    arrow(ax, rx_mid, 6.32, rx_mid, 5.12, lw=0.95)
    ax.text(8.22, 6.52, "FTQC", fontsize=6.1, fontweight="bold", ha="left")

    draw_box(ax, 5.70, 3.20, 4.10, 1.90,
             "4b. Fault-Tolerant\nLogical Execution",
             subtext="Surface-code logical qubits;\nMagic-state T-gate distillation;\nReal-time syndrome decoding",
             title_size=6.1, fill="#EDEDED", lw=0.95, rounding=0.05, subsize=5.40,
             title_y_frac=0.77, sub_y_frac=0.31)

    ax.plot([lx_mid, lx_mid], [3.20, 2.80], lw=0.95, color="black")
    ax.plot([rx_mid, rx_mid], [3.20, 2.80], lw=0.95, color="black")
    ax.plot([lx_mid, rx_mid], [2.80, 2.80], lw=0.95, color="black")
    arrow(ax, cx, 2.80, cx, 2.38, lw=0.95)

    draw_box(ax, bx, 1.38, bw, 0.98,
             "5. Shot Readout, Estimation & Parameter Update",
             subtext="Bitstring Sampling (M shots), Expectation ⟨Ô⟩ ± σ/√M & Classical Update",
             title_size=6.2, fill="#FFFFFF", lw=1.0, rounding=0.06, subsize=5.40)
    arrow(ax, cx, 1.38, cx, 0.92, lw=0.95)

    draw_box(ax, bx, 0.10, bw, 0.80,
             "6. Benchmarking Against Classical Baselines",
             subtext="Accuracy, Runtime, Shot Overhead, Noise Resilience & Scaling Audit",
             title_size=6.2, fill="#EAEAEA", lw=1.0, rounding=0.08, subsize=5.40)

    ax.plot([bx, 0.65], [1.87, 1.87], ls="--", lw=0.90, color="black")
    ax.plot([0.65, 0.65], [1.87, 8.30], ls="--", lw=0.90, color="black")
    arrow(ax, 0.65, 8.30, bx, 8.30, lw=0.90, ls="--")
    ax.text(0.30, 5.10, "Iterative Variational / Transpiler Loop", fontsize=5.8, rotation=90,
            ha="center", va="center")

    save_all_raster_and_vector(fig, 2)


# ==============================================================================
# FIGURE 3: HYBRID QUANTUM-CLASSICAL INTELLIGENT COMPUTING ARCHITECTURE
# ==============================================================================
def build_figure_3():
    fig, ax = plt.subplots(figsize=(4.40, 4.55), dpi=600)
    fig.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    ax.set_xlim(0.0, 10.0)
    ax.set_ylim(0.0, 10.0)
    ax.axis("off")

    qpu_box = FancyBboxPatch((0.18, 5.35), 9.64, 4.35, boxstyle="round,pad=0.02,rounding_size=0.08",
                             facecolor="#FAFAFA", edgecolor="black", lw=1.10, linestyle="--", zorder=1)
    ax.add_patch(qpu_box)
    ax.text(5.00, 9.28, r"QUANTUM PROCESSING UNIT (QPU) — HILBERT SPACE $\mathcal{H}^{\otimes n}$",
            fontsize=6.4, fontweight="bold", ha="center", va="center")

    cpu_box = FancyBboxPatch((0.18, 0.25), 9.64, 3.85, boxstyle="round,pad=0.02,rounding_size=0.08",
                             facecolor="#FFFFFF", edgecolor="black", lw=1.10, linestyle="-", zorder=1)
    ax.add_patch(cpu_box)
    ax.text(5.00, 0.62, "CLASSICAL CO-PROCESSOR (CPU / GPU) — PARAMETER & DATA SPACE",
            fontsize=6.3, fontweight="bold", ha="center", va="center")

    ax.text(0.82, 7.18, r"$|0\rangle^{\otimes n}$", fontsize=7.4, fontweight="bold", ha="center", va="center")
    arrow(ax, 1.28, 7.18, 1.72, 7.18, lw=0.95)

    draw_box(ax, 1.72, 6.00, 2.58, 2.36,
             "Quantum Feature\nMap Encoding",
             subtext=r"$U_\Phi(\tilde{\mathbf{x}})$" + "\nData Embedding",
             title_size=6.1, fill="#EAEAEA", lw=1.0, rounding=0.05, subsize=5.6,
             title_y_frac=0.73, sub_y_frac=0.30)
    arrow(ax, 4.30, 7.18, 4.74, 7.18, lw=0.95)

    draw_box(ax, 4.74, 6.00, 2.66, 2.36,
             "Parameterized\nAnsatz Circuit",
             subtext=r"$U_{\mathrm{var}}(\boldsymbol{\theta}^{(t)})$" + "\nTrainable Gates",
             title_size=6.1, fill="#FFFFFF", lw=1.10, rounding=0.05, subsize=5.6,
             title_y_frac=0.73, sub_y_frac=0.30)
    arrow(ax, 7.40, 7.18, 7.82, 7.18, lw=0.95)

    draw_box(ax, 7.82, 6.00, 1.80, 2.36,
             "Projective\nReadout",
             subtext=r"$\{\hat{O}_k\}_{k=1}^J$" + "\nM Shots",
             title_size=6.1, fill="#F2F2F2", lw=1.0, rounding=0.05, subsize=5.6,
             title_y_frac=0.73, sub_y_frac=0.30)

    draw_box(ax, 0.44, 1.12, 2.78, 2.60,
             "Classical Data\nPreprocessing",
             subtext=r"$\mathbf{x} \in \mathbb{R}^d \mapsto \tilde{\mathbf{x}}$" + "\nScaling / PCA",
             title_size=6.1, fill="#F4F4F4", lw=0.95, rounding=0.05, subsize=5.6,
             title_y_frac=0.74, sub_y_frac=0.30)

    draw_box(ax, 3.66, 1.12, 2.84, 2.60,
             "Optimizer &\nGradient Engine",
             subtext=r"Param.-Shift $\nabla_{\boldsymbol{\theta}}\mathcal{L}$" + "\nAdam / SPSA",
             title_size=6.1, fill="#FFFFFF", lw=1.0, rounding=0.05, subsize=5.6,
             title_y_frac=0.74, sub_y_frac=0.30)

    draw_box(ax, 6.94, 1.12, 2.66, 2.60,
             "Empirical Cost &\nDecision Output",
             subtext=r"$\mathcal{L}(\boldsymbol{\theta}, \{\langle\hat{O}_k\rangle\})$" + "\nPrediction " + r"$\hat{y}$",
             title_size=6.1, fill="#EAEAEA", lw=0.95, rounding=0.05, subsize=5.6,
             title_y_frac=0.74, sub_y_frac=0.30)

    arrow(ax, 2.55, 3.72, 2.55, 6.00, lw=1.0)
    ax.text(2.40, 4.74, r"Encoded $\tilde{\mathbf{x}}$", fontsize=6.0, ha="right", va="center")

    arrow(ax, 8.65, 6.00, 8.65, 3.72, lw=1.0)
    ax.text(8.50, 4.74, r"$\hat{\mu}_k \pm \sigma/\sqrt{M}$", fontsize=6.0, ha="right", va="center")

    arrow(ax, 6.94, 2.42, 6.50, 2.42, lw=0.95)

    arrow(ax, 5.55, 3.72, 5.55, 6.00, lw=1.05, ls="--")
    ax.text(5.40, 4.74, r"Update $\boldsymbol{\theta}^{(t+1)}$", fontsize=6.0, ha="right", va="center")

    save_all_raster_and_vector(fig, 3)


# ==============================================================================
# FIGURE 4: COMPARATIVE SCALING PLOT + REGIME CLASSIFICATION MATRIX
# ==============================================================================
def build_figure_4():
    fig = plt.figure(figsize=(4.40, 5.40), dpi=600)
    gs = fig.add_gridspec(2, 1, height_ratios=[1.12, 1.02], hspace=0.30,
                          left=0.14, right=0.96, top=0.94, bottom=0.03)

    ax1 = fig.add_subplot(gs[0, 0])
    n = np.array([10, 20, 30, 40, 50, 60, 70, 80, 90, 100], dtype=float)
    log10_classical_search = n * np.log10(2.0)
    log10_grover = 0.5 * n * np.log10(2.0) + np.log10(np.pi / 4.0)
    lnN = n * np.log(2.0)
    c_gnfs = (64.0 / 9.0) ** (1.0 / 3.0)
    log10_gnfs = (c_gnfs * (lnN ** (1.0 / 3.0)) * ((np.log(lnN)) ** (2.0 / 3.0))) / np.log(10.0)
    log10_shor = 3.0 * np.log10(n)

    ax1.plot(n, log10_classical_search, color="black", lw=1.20, ls="-", marker="o", ms=3.5,
             label=r"Classical Search $O(2^n)$")
    ax1.plot(n, log10_grover, color="#222222", lw=1.15, ls="--", marker="s", ms=3.4,
             markerfacecolor="white", label=r"Grover Search $O(2^{n/2})$")
    ax1.plot(n, log10_gnfs, color="#333333", lw=1.15, ls="-.", marker="^", ms=3.8,
             label=r"Classical GNFS $L_N[1/3, c]$")
    ax1.plot(n, log10_shor, color="black", lw=1.20, ls=":", marker="D", ms=3.4,
             markerfacecolor="white", label=r"Shor Factoring $O(n^3)$")

    ax1.set_title("(a) Asymptotic Query / Gate Scaling by Problem Bit-Width",
                  fontsize=7.2, fontweight="bold", loc="center", pad=6)
    ax1.set_xlabel(r"Problem Bit-Width $n = \log_2 N$ (qubits / classical bits)", fontsize=7.0)
    ax1.set_ylabel(r"$\log_{10}(\mathrm{Operations})$", fontsize=7.0)
    ax1.set_xlim(8, 103)
    ax1.set_ylim(0, 33)
    ax1.grid(True, linestyle=":", linewidth=0.5, color="#AAAAAA")
    ax1.tick_params(labelsize=6.8)
    ax1.legend(loc="upper left", fontsize=6.3, frameon=True, facecolor="white",
               edgecolor="black", framealpha=0.96)

    ax2 = fig.add_subplot(gs[1, 0])
    ax2.set_position([0.03, 0.02, 0.94, 0.42])
    ax2.set_xlim(0.0, 10.0)
    ax2.set_ylim(0.0, 6.4)
    ax2.axis("off")
    ax2.text(5.0, 6.08, "(b) Algorithmic Regime-Selection Matrix",
             fontsize=7.2, fontweight="bold", ha="center")

    ax2.text(3.38, 5.42, "Shallow Circuit Depth\n(NISQ + Error Mitigation)",
             fontsize=6.5, fontweight="bold", ha="center", va="center")
    ax2.text(7.62, 5.42, "Deep Coherent Depth\n(Fault-Tolerant QEC)",
             fontsize=6.5, fontweight="bold", ha="center", va="center")
    ax2.text(0.58, 3.82, "Structured /\nQuantum Data", fontsize=6.2, fontweight="bold",
             ha="center", va="center", rotation=90)
    ax2.text(0.58, 1.38, "Generic\nClassical Data", fontsize=6.2, fontweight="bold",
             ha="center", va="center", rotation=90)

    draw_box(ax2, 1.25, 2.72, 4.22, 2.20,
             "Pre-FT Utility Regime",
             subtext="Symmetry-preserving VQE,\nClassical shadows & local\nquantum spin dynamics",
             title_size=6.5, fill="#F0F0F0", lw=0.95, rounding=0.04, subsize=5.8,
             title_y_frac=0.74, sub_y_frac=0.31)

    draw_box(ax2, 5.55, 2.72, 4.22, 2.20,
             "Provable Speedup Regime",
             subtext="Shor / Discrete-Log Kernels,\nQPE Chemistry & sparse HHL\nwith efficient state prep",
             title_size=6.5, fill="#FFFFFF", lw=1.10, rounding=0.04, subsize=5.8,
             title_y_frac=0.74, sub_y_frac=0.31)

    draw_box(ax2, 1.25, 0.28, 4.22, 2.20,
             "Heuristic Benchmark Regime",
             subtext="QAOA combinatorial opt. &\nshallow QNNs; requires audit\nagainst classical ML/MIP",
             title_size=6.5, fill="#FFFFFF", lw=0.95, rounding=0.04, subsize=5.8,
             title_y_frac=0.74, sub_y_frac=0.31)

    draw_box(ax2, 5.55, 0.28, 4.22, 2.20,
             "Quadratic / Conditional",
             subtext="Grover / Amplitude Amplif.\nO(√N); I/O & QRAM overheads\nconstrain net speedup",
             title_size=6.5, fill="#EAEAEA", lw=0.95, rounding=0.04, subsize=5.8,
             title_y_frac=0.74, sub_y_frac=0.31)

    save_all_raster_and_vector(fig, 4)


# ==============================================================================
# NATIVE EDITABLE POWERPOINT VECTOR SOURCE FILES (.pptx) FOR FIGURES 1–4
# ==============================================================================
def add_ppt_box(slide, left, top, width, height, text, font_size=8.0, bold=False,
                fill_gray=255, shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(
        shape_type, PptInches(left), PptInches(top), PptInches(width), PptInches(height)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = PptRGB(fill_gray, fill_gray, fill_gray)
    sh.line.color.rgb = PptRGB(0, 0, 0)
    sh.line.width = PptPt(1.0)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = PptInches(0.05)
    tf.margin_right = PptInches(0.05)
    tf.margin_top = PptInches(0.03)
    tf.margin_bottom = PptInches(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Times New Roman"
    r.font.size = PptPt(font_size)
    r.font.bold = bold
    r.font.color.rgb = PptRGB(0, 0, 0)
    return sh


def add_ppt_line(slide, x1, y1, x2, y2, width_pt=1.0):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, PptInches(x1), PptInches(y1), PptInches(x2), PptInches(y2)
    )
    conn.line.color.rgb = PptRGB(0, 0, 0)
    conn.line.width = PptPt(width_pt)
    return conn


def add_ppt_label(slide, left, top, width, height, text, font_size=7.8, bold=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Times New Roman"
    r.font.size = PptPt(font_size)
    r.font.bold = bold
    r.font.color.rgb = PptRGB(0, 0, 0)
    return tb


def build_editable_pptx_sources():
    prs1 = Presentation()
    prs1.slide_width = PptInches(4.50)
    prs1.slide_height = PptInches(5.40)
    s1 = prs1.slides.add_slide(prs1.slide_layouts[6])
    add_ppt_label(s1, 0.08, 0.08, 1.85, 0.25, "(a) Qubit Bloch Sphere", 8.0, True)
    add_ppt_box(s1, 0.28, 0.45, 1.35, 1.35, "", fill_gray=255, shape_type=MSO_SHAPE.OVAL)
    add_ppt_box(s1, 0.28, 0.95, 1.35, 0.35, "", fill_gray=255, shape_type=MSO_SHAPE.OVAL)
    add_ppt_line(s1, 0.955, 0.35, 0.955, 1.92, 1.0)
    add_ppt_line(s1, 0.955, 1.125, 1.42, 0.66, 1.5)
    add_ppt_label(s1, 0.72, 0.16, 0.55, 0.20, "|0⟩", 7.8, True)
    add_ppt_label(s1, 0.72, 1.90, 0.55, 0.20, "|1⟩", 7.8, True)
    add_ppt_label(s1, 1.36, 0.46, 0.55, 0.22, "|ψ⟩", 8.0, True)

    add_ppt_label(s1, 1.98, 0.08, 2.40, 0.25, "(b) Canonical 3-Qubit Circuit", 8.0, True)
    for idx, yw in enumerate([0.55, 1.00, 1.45]):
        add_ppt_label(s1, 1.92, yw - 0.12, 0.38, 0.24, f"|0⟩q{idx}", 7.2)
        add_ppt_line(s1, 2.30, yw, 4.30, yw, 1.0)
    add_ppt_line(s1, 2.30, 1.88, 4.35, 1.88, 1.2)
    add_ppt_label(s1, 1.92, 1.76, 0.38, 0.24, "c3", 7.2)
    add_ppt_box(s1, 2.42, 0.40, 0.32, 0.30, "H", 7.8, True, 255, MSO_SHAPE.RECTANGLE)
    add_ppt_box(s1, 2.40, 1.30, 0.52, 0.30, "Ry(θ1)", 7.0, False, 240, MSO_SHAPE.RECTANGLE)
    add_ppt_line(s1, 2.98, 0.55, 2.98, 1.00, 1.1)
    add_ppt_box(s1, 2.90, 0.92, 0.16, 0.16, "+", 7.2, True, 255, MSO_SHAPE.OVAL)
    add_ppt_line(s1, 3.24, 1.00, 3.24, 1.45, 1.1)
    add_ppt_box(s1, 3.16, 1.37, 0.16, 0.16, "+", 7.2, True, 255, MSO_SHAPE.OVAL)
    add_ppt_box(s1, 3.40, 0.85, 0.52, 0.30, "Rz(θ2)", 7.0, False, 240, MSO_SHAPE.RECTANGLE)
    for yw in [0.55, 1.00, 1.45]:
        add_ppt_box(s1, 4.02, yw - 0.14, 0.30, 0.28, "M", 7.2, True, 255, MSO_SHAPE.RECTANGLE)

    add_ppt_label(s1, 0.10, 2.22, 4.30, 0.26, "(c) Four-Tier Quantum Execution Architecture", 8.0, True)
    stack_items = [
        (2.55, "Tier 4: Quantum Algorithms & Intelligent Computing (VQE, QAOA, QML, Shor, Grover)", 255),
        (3.25, "Tier 3: Logical Compilation, Gate Decomposition & Hardware Topology Routing", 242),
        (3.95, "Tier 2: Error Mitigation (NISQ ZNE/PEC) & Fault-Tolerant QEC (Surface Codes)", 255),
        (4.65, "Tier 1: Physical Qubit Hardware (Superconducting, Trapped-Ion, Neutral-Atom)", 232),
    ]
    for yt, txt, gr in stack_items:
        add_ppt_box(s1, 0.16, yt, 4.18, 0.54, txt, 7.2, False, gr)

    prs2 = Presentation()
    prs2.slide_width = PptInches(4.50)
    prs2.slide_height = PptInches(5.40)
    s2 = prs2.slides.add_slide(prs2.slide_layouts[6])
    wf_steps = [
        (0.14, "1. Problem Specification: Cost Hamiltonian H_C, Oracle f(x), or Dataset {(x_i, y_i)}", 235),
        (0.84, "2. Algorithmic Paradigm & Data Encoding Design (Amplitude / Angle / Kernel Map)", 255),
        (1.54, "3. Hardware-Aware Compilation: Native Gate Synthesis & Qubit Routing", 244),
    ]
    for yt, txt, gr in wf_steps:
        add_ppt_box(s2, 0.32, yt, 3.86, 0.52, txt, 7.2, False, gr)
    add_ppt_box(s2, 1.15, 2.24, 2.20, 0.62, "Hardware Regime?\nNISQ vs. Fault-Tolerant", 7.2, True, 255, MSO_SHAPE.DIAMOND)
    add_ppt_box(s2, 0.18, 3.04, 1.98, 0.90, "4a. NISQ Execution:\nShallow PQC + Error Mitigation\n(ZNE, PEC, Readout Mitigation)", 6.8, False, 255)
    add_ppt_box(s2, 2.34, 3.04, 1.98, 0.90, "4b. Fault-Tolerant Execution:\nLogical Qubits (p < p_thr),\nSurface Code & T-Distillation", 6.8, False, 235)
    add_ppt_box(s2, 0.32, 4.14, 3.86, 0.50, "5. Shot Readout, Observable Estimation ⟨Ô⟩ ± σ/√M & Optimizer Update", 7.2, False, 255)
    add_ppt_box(s2, 0.32, 4.78, 3.86, 0.48, "6. Benchmarking Against State-of-the-Art Classical Baselines", 7.2, True, 232)

    prs3 = Presentation()
    prs3.slide_width = PptInches(4.50)
    prs3.slide_height = PptInches(4.55)
    s3 = prs3.slides.add_slide(prs3.slide_layouts[6])
    add_ppt_box(s3, 0.12, 0.14, 4.26, 1.95, "", 8.0, False, 250)
    add_ppt_label(s3, 0.18, 0.18, 4.14, 0.28, "QUANTUM PROCESSING UNIT (QPU) — HILBERT SPACE", 7.4, True)
    add_ppt_box(s3, 0.25, 0.62, 1.25, 1.25, "Quantum Feature\nMap Encoding\nU_Φ(x)", 7.2, True, 234)
    add_ppt_box(s3, 1.66, 0.62, 1.32, 1.25, "Variational\nAnsatz Circuit\nU_var(θ^(t))", 7.2, True, 255)
    add_ppt_box(s3, 3.14, 0.62, 1.08, 1.25, "Projective\nReadout\n{Ô_k}", 7.2, True, 242)

    add_ppt_box(s3, 0.12, 2.40, 4.26, 1.98, "", 8.0, False, 255)
    add_ppt_label(s3, 0.18, 2.44, 4.14, 0.28, "CLASSICAL CO-PROCESSOR (CPU / GPU)", 7.4, True)
    add_ppt_box(s3, 0.25, 2.88, 1.25, 1.30, "Classical Data\nPreprocessing\nx ∈ R^d → x̃", 7.0, False, 242)
    add_ppt_box(s3, 1.66, 2.88, 1.32, 1.30, "Optimizer &\nGradient Engine\nAdam / SPSA", 7.0, True, 255)
    add_ppt_box(s3, 3.14, 2.88, 1.08, 1.30, "Empirical Loss\nL(θ) & Output\nPrediction ŷ", 7.0, False, 234)
    add_ppt_line(s3, 0.88, 2.88, 0.88, 1.87, 1.2)
    add_ppt_line(s3, 3.68, 1.87, 3.68, 2.88, 1.2)
    add_ppt_line(s3, 2.32, 2.88, 2.32, 1.87, 1.2)

    prs4 = Presentation()
    prs4.slide_width = PptInches(4.50)
    prs4.slide_height = PptInches(5.40)
    s4 = prs4.slides.add_slide(prs4.slide_layouts[6])
    add_ppt_label(s4, 0.14, 0.08, 4.22, 0.28, "(a) Asymptotic Complexity Scaling Comparison", 7.8, True)
    add_ppt_box(s4, 0.45, 0.40, 3.85, 2.05, "", 8.0, False, 255, MSO_SHAPE.RECTANGLE)
    n_vals = [10, 30, 50, 70, 100]
    def px(nv): return 0.45 + ((nv - 10) / 90.0) * 3.85
    def py(lv): return 2.45 - (min(lv, 31.0) / 31.0) * 2.05
    for fn, wpt in [
        (lambda nv: nv * math.log10(2), 1.4),
        (lambda nv: 0.5 * nv * math.log10(2), 1.1),
        (lambda nv: ((64/9)**(1/3) * (nv*math.log(2))**(1/3) * (math.log(nv*math.log(2)))**(2/3)) / math.log(10), 1.1),
        (lambda nv: 3.0 * math.log10(nv), 1.3),
    ]:
        for i in range(len(n_vals) - 1):
            add_ppt_line(s4, px(n_vals[i]), py(fn(n_vals[i])), px(n_vals[i+1]), py(fn(n_vals[i+1])), wpt)
    add_ppt_label(s4, 0.52, 0.46, 2.35, 0.72,
                  "Classical Search O(2^n)  |  Grover O(2^(n/2))\nClassical GNFS L_N[1/3, c]  |  Shor O(n^3)",
                  6.8, False, PP_ALIGN.LEFT)
    add_ppt_label(s4, 0.45, 2.48, 3.85, 0.24, "Register Bit-Width n = log2(N)   [Vertical: log10(Operations)]", 7.0, False)

    add_ppt_label(s4, 0.14, 2.82, 4.22, 0.28, "(b) Algorithmic Regime-Selection Matrix", 7.8, True)
    add_ppt_box(s4, 0.18, 3.16, 2.02, 1.02, "Pre-FT Utility (Structured + Shallow):\nSymmetry VQE, Classical Shadows,\nLocal Quantum Simulation", 6.8, False, 240)
    add_ppt_box(s4, 2.30, 3.16, 2.02, 1.02, "Provable Speedup (Structured + Deep):\nShor / Discrete-Log Kernels,\nQPE Chemistry, Sparse HHL", 6.8, True, 255)
    add_ppt_box(s4, 0.18, 4.26, 2.02, 1.02, "Heuristic Benchmark (Generic + Shallow):\nQAOA & Shallow QNNs;\nRequires Classical Baseline Audit", 6.8, False, 255)
    add_ppt_box(s4, 2.30, 4.26, 2.02, 1.02, "Quadratic / Conditional (Generic + Deep):\nGrover O(√N); QRAM & State-Prep\nOverheads Constrain Speedup", 6.8, False, 232)

    for idx, prs in enumerate([prs1, prs2, prs3, prs4], 1):
        for out_d in OUT_DIRS:
            prs.save(os.path.join(out_d, f"Figure_{idx}.pptx"))


if __name__ == "__main__":
    build_figure_1()
    build_figure_2()
    build_figure_3()
    build_figure_4()
    build_editable_pptx_sources()
    print("Regenerated all 4 figures cleanly!")
